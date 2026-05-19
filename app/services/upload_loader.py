"""UI'dan yüklenen dosyaları ham payload (canonical veya tabular) olarak okur."""

from __future__ import annotations

import csv
import json
import os
from pathlib import Path
from typing import Any

_CANONICAL_KEYS = ("company_profile", "invoices", "suppliers", "documents")
_JSON_EXTENSIONS = {".json"}
_CSV_EXTENSIONS = {".csv"}
_EXCEL_EXTENSIONS = {".xlsx", ".xls"}
_DOCUMENT_EXTENSIONS = {".pdf", ".png", ".jpg", ".jpeg", ".webp"}
_SUPPORTED_EXTENSIONS = (
    _JSON_EXTENSIONS | _CSV_EXTENSIONS | _EXCEL_EXTENSIONS | _DOCUMENT_EXTENSIONS
)

# Internal metadata key on each tabular row (Excel 1-based row number)
EXCEL_ROW_META = "_excel_row"


def get_max_upload_rows() -> int:
    return int(os.getenv("MAX_UPLOAD_ROWS", "50000"))


def get_large_excel_threshold() -> int:
    """Bu satır sayısının üzerinde openpyxl read_only okuma kullanılır."""
    return int(os.getenv("LARGE_EXCEL_READ_ONLY_THRESHOLD", "10000"))


def strip_row_metadata(row: dict[str, Any]) -> dict[str, Any]:
    """Export / Excel yazımı öncesi internal satır meta anahtarını kaldırır."""
    return {k: v for k, v in row.items() if k != EXCEL_ROW_META}


def is_document_path(file_path: str) -> bool:
    """PDF veya görsel belge yolu mu?"""
    return Path(file_path).suffix.lower() in _DOCUMENT_EXTENSIONS


def is_tabular_or_json_path(file_path: str) -> bool:
    """Tabular veya kanonik JSON yükleme modu."""
    suffix = Path(file_path).suffix.lower()
    return suffix in (_JSON_EXTENSIONS | _CSV_EXTENSIONS | _EXCEL_EXTENSIONS)


def load_uploaded_file(file_path: str) -> dict[str, Any]:
    """Dosya uzantısına göre canonical, tabular veya belge referansı döndürür."""
    path = Path(file_path)
    if not path.is_file():
        raise RuntimeError(f"UPLOAD_FILE_NOT_FOUND: {file_path}")

    suffix = path.suffix.lower()
    if suffix in _DOCUMENT_EXTENSIONS:
        return {"mode": "document", "path": str(path.resolve())}
    if suffix in _JSON_EXTENSIONS:
        return _load_json(path)
    if suffix in _CSV_EXTENSIONS:
        return _load_tabular(path, source_name=path.name)
    if suffix in _EXCEL_EXTENSIONS:
        return _load_tabular(path, source_name=path.name)

    raise RuntimeError(f"UPLOAD_UNSUPPORTED_FORMAT: {suffix}")


def _load_json(path: Path) -> dict[str, Any]:
    try:
        with path.open(encoding="utf-8") as file:
            data = json.load(file)
    except json.JSONDecodeError as exc:
        raise RuntimeError(f"UPLOAD_JSON_PARSE_ERROR: {exc}") from exc

    if not isinstance(data, dict):
        raise RuntimeError("UPLOAD_JSON_PARSE_ERROR: kök nesne dict olmalı")

    if "canonical" in data and isinstance(data["canonical"], dict):
        canonical = _normalize_canonical(data["canonical"])
    else:
        canonical = _normalize_canonical(data)

    return {"mode": "canonical", "canonical": canonical}


def _normalize_canonical(data: dict[str, Any]) -> dict[str, Any]:
    from app.services.verification.inventory_trust_policy import normalize_canonical_document

    raw_docs = data.get("documents") if isinstance(data.get("documents"), list) else []
    documents = [
        normalize_canonical_document(d) if isinstance(d, dict) else d
        for d in raw_docs
    ]
    return {
        "company_profile": data.get("company_profile") or {},
        "invoices": data.get("invoices") if isinstance(data.get("invoices"), list) else [],
        "suppliers": data.get("suppliers") if isinstance(data.get("suppliers"), list) else [],
        "documents": documents,
    }


def _load_tabular(path: Path, source_name: str) -> dict[str, Any]:
    try:
        rows, columns = _read_tabular_rows(path)
    except RuntimeError:
        raise
    except Exception as exc:
        raise RuntimeError(f"UPLOAD_TABULAR_PARSE_ERROR: {exc}") from exc

    max_rows = get_max_upload_rows()
    if len(rows) > max_rows:
        raise RuntimeError(
            f"UPLOAD_ROW_LIMIT_EXCEEDED: Dosyada {len(rows)} veri satırı var; "
            f"üst sınır {max_rows}. Lütfen dönemi bölün veya filtreleyin."
        )

    return {
        "mode": "tabular",
        "tabular": {
            "rows": rows,
            "columns": columns,
            "source_name": source_name,
        },
    }


def _read_tabular_rows(path: Path) -> tuple[list[dict[str, Any]], list[str]]:
    suffix = path.suffix.lower()

    try:
        import pandas as pd
    except ImportError:
        pd = None  # type: ignore[assignment]

    if suffix == ".csv":
        if pd is not None:
            frame = pd.read_csv(path)
            return _frame_to_rows(frame, pd)
        return _read_csv_fallback(path)

    if suffix in _EXCEL_EXTENSIONS:
        if pd is None:
            raise RuntimeError(
                "pandas gerekli — pip install pandas openpyxl"
            )
        if suffix == ".xlsx" and _should_use_read_only_excel(path):
            return _read_excel_openpyxl(path)
        frame = pd.read_excel(path, engine="openpyxl" if suffix == ".xlsx" else None)
        return _frame_to_rows(frame, pd)

    raise RuntimeError(f"UPLOAD_UNSUPPORTED_FORMAT: {suffix}")


def _should_use_read_only_excel(path: Path) -> bool:
    """Büyük xlsx dosyalarında bellek dostu okuma."""
    threshold = get_large_excel_threshold()
    try:
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        max_row = ws.max_row or 0
        wb.close()
        # başlık + veri satırları
        return max_row > threshold + 1
    except Exception:
        return False


def _read_excel_openpyxl(path: Path) -> tuple[list[dict[str, Any]], list[str]]:
    from openpyxl import load_workbook

    max_rows = get_max_upload_rows()
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.active
        row_iter = ws.iter_rows(values_only=True)
        header = next(row_iter, None)
        if not header:
            return [], []

        columns = [
            str(cell).strip() if cell is not None else f"Sütun{i + 1}"
            for i, cell in enumerate(header)
        ]
        rows: list[dict[str, Any]] = []
        excel_row = 2
        for values in row_iter:
            if len(rows) >= max_rows:
                raise RuntimeError(
                    f"UPLOAD_ROW_LIMIT_EXCEEDED: üst sınır {max_rows} veri satırı"
                )
            if values is None or all(v is None or str(v).strip() == "" for v in values):
                excel_row += 1
                continue
            row: dict[str, Any] = {EXCEL_ROW_META: excel_row}
            for col_name, val in zip(columns, values):
                if val is None or (isinstance(val, str) and not val.strip()):
                    row[col_name] = None
                else:
                    row[col_name] = val
            rows.append(row)
            excel_row += 1
        return rows, columns
    finally:
        wb.close()


def _frame_to_rows(frame: Any, pd: Any) -> tuple[list[dict[str, Any]], list[str]]:
    columns = [str(column) for column in frame.columns.tolist()]
    rows: list[dict[str, Any]] = []
    for row_idx, record in enumerate(frame.to_dict(orient="records")):
        row: dict[str, Any] = {EXCEL_ROW_META: row_idx + 2}
        for key, value in record.items():
            if pd.isna(value):
                row[str(key)] = None
            else:
                row[str(key)] = value
        rows.append(row)
    return rows, columns


def _read_csv_fallback(path: Path) -> tuple[list[dict[str, Any]], list[str]]:
    with path.open(encoding="utf-8-sig", newline="") as file:
        reader = csv.DictReader(file)
        columns = list(reader.fieldnames or [])
        rows: list[dict[str, Any]] = []
        for row_idx, raw in enumerate(reader):
            row = dict(raw)
            row[EXCEL_ROW_META] = row_idx + 2
            rows.append(row)
    return rows, columns
